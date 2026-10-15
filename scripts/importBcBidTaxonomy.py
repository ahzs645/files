#!/usr/bin/env python3
"""Import a versioned UNSPSC taxonomy into the EXISTING primary Procurement catalog.

Dry-run by default; --apply publishes only after full node/change read-back.
No database creation, SQL writes, deletions, or historical-version guessing.
Requires read-only openpyxl. Existing notices/files/checkpoints are preserved.
"""
import argparse
import collections
import hashlib
import json
import pathlib
import urllib.request
import urllib.parse
import http.client
import time
from backfillBcBidMetadata import SOURCE, SHEET, workbook_lookup, source_categories, enrich

VERSION = 'UNv260801'
ACTIVE = 'procurement:taxonomy:active:unspsc'
RELEASE = 'taxonomy-release:UNSPSC:' + VERSION


def encoded(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':')).encode()


def term_id(code):
    return 'taxonomy-term:UNSPSC:' + VERSION + ':' + code


def workbook_records(path):
    import openpyxl
    lookup = workbook_lookup(path)  # Header/code validation before any writes.
    nodes = {}
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for number, row in enumerate(workbook[SHEET].iter_rows(min_row=8, values_only=True), 8):
        ancestors = []
        for index, level in [(1, 'segment'), (3, 'family'), (5, 'class'), (7, 'commodity')]:
            if not row[index] or not row[index + 1]:
                raise ValueError('Incomplete source hierarchy.')
            code, title = str(int(row[index])), str(row[index + 1]).strip()
            parent = term_id(ancestors[-1]['code']) if ancestors else None
            data = {'schemaVersion': 1, 'scheme': 'UNSPSC', 'version': VERSION, 'releaseId': RELEASE,
                    'code': code, 'title': title, 'level': level, 'parentId': parent,
                    'ancestors': list(ancestors), 'path': '|' + '|'.join([a['code'] for a in ancestors] + [code]) + '|',
                    'definition': str(row[9]).strip() if index == 7 and row[9] else None, 'sourceRow': number}
            if code in nodes:
                old = nodes[code]['data']
                if any(old[key] != data[key] for key in ['title', 'level', 'parentId', 'ancestors', 'definition']):
                    raise ValueError('Conflicting source taxonomy node: ' + code)
            else:
                nodes[code] = {'id': term_id(code), 'kind': 'taxonomy-term', 'title': title, 'data': data}
            ancestors.append({'code': code, 'title': title, 'level': level})
    identifiers = {term_id(code) for code in nodes}
    for node in nodes.values():
        if node['data']['parentId'] and node['data']['parentId'] not in identifiers:
            raise ValueError('Missing taxonomy parent.')
    changes = []
    for name in ["Whats' New", "What's Changed", "What's Deleted"]:
        rows = workbook[name].iter_rows(values_only=True)
        headers = next(rows)
        for number, row in enumerate(rows, 2):
            if not row[0]:
                continue
            data = dict(zip(headers, row))
            data.update({'schemaVersion': 1, 'scheme': 'UNSPSC', 'releaseId': RELEASE, 'sourceSheet': name, 'sourceRow': number})
            changes.append({'id': f'taxonomy-change:UNSPSC:{VERSION}:{len(changes):04d}', 'kind': 'taxonomy-change',
                            'title': str(data.get('effective_title') or data.get('changed_title')), 'data': data})
    workbook.close()
    counts = dict(collections.Counter(node['data']['level'] for node in nodes.values()))
    provenance = {'system': 'UNSPSC', 'version': VERSION, 'sourceUrl': SOURCE,
                  'sourceSha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheet': SHEET,
                  'method': 'unique-exact-title-match'}
    ordered = [nodes[code] for code in sorted(nodes)]
    release = {'id': RELEASE, 'kind': 'taxonomy-release', 'title': 'UNSPSC ' + VERSION,
               'data': {'schemaVersion': 1, 'scheme': 'UNSPSC', 'version': VERSION, 'status': 'ready',
                        'previousVersion': 'UNv250901', 'previousReleaseCoverage': 'change-log-only',
                        'counts': counts, 'terms': len(ordered), 'changes': len(changes), 'provenance': provenance,
                        'termsSha256': hashlib.sha256(encoded(ordered)).hexdigest(),
                        'changesSha256': hashlib.sha256(encoded(changes)).hexdigest()}}
    return ordered, changes, release, lookup


def index_entries(lookup):
    shards = {format(index, '02x'): [] for index in range(256)}
    for title in sorted(lookup):
        shard = hashlib.sha256(title.encode()).hexdigest()[:2]
        shards[shard].append({'title': title, 'termIds': [term_id(code) for code in sorted(lookup[title])]})
    return [{'key': f'procurement:taxonomy:index:{VERSION}:{key}',
             'value': {'schemaVersion': 1, 'releaseId': RELEASE, 'items': items}}
            for key, items in shards.items()]


def linked_notice(record, lookup, release, categories=None, details=None):
    original = record['data']
    if original.get('historical') or original.get('commodityScheme') == 'GSIN' or (original.get('type') or '').startswith('Historic Record') or any('GSIN' in f.get('label', '') for f in original.get('detailFields', [])):
        return record, []
    observed = (categories or {}).get(str(original.get('processId') or original.get('sourceKey')))
    applicable = categories or {}
    if observed and (original.get('sourceCapturedAt') or '') > observed['capturedAt']:
        applicable = {}
    data = enrich(original, lookup, release['data']['provenance'], applicable, details)
    names = data.get('commodities') or []
    main = data.get('mainCommodity') if data.get('mainCommodity') in names else None
    fingerprint = hashlib.sha256(encoded([RELEASE, names, main])).hexdigest()
    data.update({'mainCommodity': main, 'classificationLinkSet': fingerprint, 'taxonomyReleaseId': RELEASE})
    links = []
    for item in data['commodityClassifications']:
        role = 'main' if item['sourceLabel'] == main else 'other' if main else 'unspecified'
        identifier = f"opportunity-classification:{record['id']}:UNSPSC:{VERSION}:{fingerprint}:{item['code']}"
        if any(link['id'] == identifier for link in links):
            continue
        links.append({'id': identifier, 'kind': 'opportunity-classification', 'title': item['sourceLabel'],
                      'data': {'schemaVersion': 1, 'recordId': record['id'], 'termId': term_id(item['code']),
                               'releaseId': RELEASE, 'scheme': 'UNSPSC', 'version': VERSION, 'code': item['code'],
                               'role': role, 'rawLabel': item['sourceLabel'], 'inputFingerprint': fingerprint,
                               'method': 'unique-exact-title-match'}})
    return {'id': record['id'], 'kind': record['kind'], 'title': record['title'], 'data': data}, links


def run(args):
    terms, changes, release, lookup = workbook_records(args.workbook)
    args.output.mkdir(parents=True, exist_ok=True)
    # Portable prepared artifacts, never hidden database storage.
    indices = index_entries(lookup)
    for name, value in [('terms.json', terms), ('changes.json', changes), ('release.json', release), ('indices.json', indices)]:
        (args.output / name).write_bytes(encoded(value))
    base = args.base.rstrip('/') + '/api/databases/catalogs/bc-bid-monitor'
    if not base.startswith('https://'):
        raise ValueError('Use the existing HTTPS catalog.')
    target = urllib.parse.urlsplit(base)
    connection = http.client.HTTPSConnection(target.hostname, target.port or 443, timeout=45)
    def call(route, body):
        nonlocal connection
        payload = encoded(body)
        for attempt in range(6):
            try:
                connection.request('POST', target.path + route, body=payload, headers={'Content-Type': 'application/json'})
                response = connection.getresponse()
                content = response.read()
                if response.status in [404, 408, 429, 500, 502, 503, 504] and attempt < 5:
                    connection.close()
                    connection = http.client.HTTPSConnection(target.hostname, target.port or 443, timeout=45)
                    time.sleep(min(2 ** attempt, 15))
                    continue
                if response.status != 200:
                    raise ValueError('Catalog API returned HTTP ' + str(response.status))
                return json.loads(content)
            except (OSError, http.client.HTTPException):
                connection.close()
                connection = http.client.HTTPSConnection(target.hostname, target.port or 443, timeout=45)
                if attempt == 5:
                    raise
                time.sleep(min(2 ** attempt, 15))
        raise ValueError('Catalog API retry limit reached.')
    def commit(rows=None, entries=None):
        for attempt in range(5):
            head = call('/read', {'ids': []})
            if not head['primary']:
                raise ValueError('Existing primary catalog required; no storage was created.')
            result = call('/commit', {'revision': head['revision'], 'records': rows or [], 'entries': entries or []})
            if not result.get('conflict'):
                return result
        raise ValueError('Repeated concurrent revisions; resume with the same workbook.')
    def save_immutable(rows):
        for offset in range(0, len(rows), 200):
            batch = rows[offset:offset + 200]
            for attempt in range(5):
                head = call('/read', {'ids': [row['id'] for row in batch]})
                if not head['primary']:
                    raise ValueError('Existing primary catalog required.')
                old = {row['id']: row for row in head['records']}
                for row in batch:
                    if row['id'] in old and any(old[row['id']][key] != row[key] for key in ['kind', 'title', 'data']):
                        raise ValueError('An immutable taxonomy entry differs: ' + row['id'])
                missing = [row for row in batch if row['id'] not in old]
                if missing:
                    result = call('/commit', {'revision': head['revision'], 'records': missing})
                    if result.get('conflict'):
                        continue
                break
            else:
                raise ValueError('Concurrent immutable import; resume safely.')
            saved = {row['id']: row for row in call('/read', {'ids': [row['id'] for row in batch]})['records']}
            for row in batch:
                if row['id'] not in saved or any(saved[row['id']][key] != row[key] for key in ['kind', 'title', 'data']):
                    raise ValueError('Taxonomy read-back mismatch: ' + row['id'])
            if offset % 4000 == 0:
                print(json.dumps({'verified': offset + len(batch), 'of': len(rows), 'kind': batch[0]['kind']}), flush=True)
    categories = source_categories(args.categories)
    details = {}
    if args.commodity_details:
        from backfillBcBidMetadata import validate_detail_capture
        details = validate_detail_capture(args.commodity_details)
    if args.opportunities:
        if args.apply:
            raise ValueError('Offline snapshots are for dry runs only; apply always merges live records.')
        records = json.loads(args.opportunities.read_text())
    else:
        records, cursor = [], ''
        for _ in range(1000):
            page = call('/list', {'kind': 'opportunity', 'after': cursor, 'limit': 100})
            if not page['primary']:
                raise ValueError('Existing primary catalog required.')
            records.extend(row for row in page['records'] if (row['data'].get('detailUrl') or '').startswith('https://bcbid.gov.bc.ca/'))
            if not page['next']:
                break
            if page['next'] <= cursor:
                raise ValueError('Catalog cursor did not advance.')
            cursor = page['next']
        else:
            raise ValueError('Opportunity inventory exceeded its bounded scan.')
    plan = {'release': release['data'], 'opportunities': len(records), 'links': 0, 'linkedOpportunities': 0, 'unresolvedLabels': []}
    unresolved = set()
    if args.apply:
        # Immutable terms are never used by the scraper until release activation.
        save_immutable(terms)
        save_immutable(changes)
        for offset in range(0, len(indices), 20):
            batch = indices[offset:offset + 20]
            for attempt in range(5):
                old = call('/workspace', {'keys': [entry['key'] for entry in batch]})
                existing = {entry['key']: entry['value'] for entry in old['entries']}
                if any(entry['key'] in existing and existing[entry['key']] != entry['value'] for entry in batch):
                    raise ValueError('An immutable taxonomy lookup differs.')
                missing = [entry for entry in batch if entry['key'] not in existing]
                if missing:
                    result = call('/commit', {'revision': old['revision'], 'entries': missing})
                    if result.get('conflict'):
                        continue
                saved = call('/workspace', {'keys': [entry['key'] for entry in batch]})
                if {entry['key']: entry['value'] for entry in saved['entries']} != {entry['key']: entry['value'] for entry in batch}:
                    raise ValueError('Taxonomy lookup read-back mismatch.')
                break
            else:
                raise ValueError('Concurrent taxonomy lookup import; resume safely.')
        save_immutable([release])
    for offset in range(0, len(records), 10):
        batch = records[offset:offset + 10]
        for attempt in range(5):
            head = call('/read', {'ids': [row['id'] for row in batch]}) if args.apply else {'records': batch}
            rows, links, baseline = [], [], []
            for old in head['records']:
                if old['kind'] != 'opportunity' or not (old['data'].get('detailUrl') or '').startswith('https://bcbid.gov.bc.ca/'):
                    raise ValueError('Opportunity scope changed.')
                row, linked = linked_notice(old, lookup, release, categories, details)
                rows.append(row); links.extend(linked); baseline.append(old)
            if args.apply:
                # Same active-set rule as the worker: staging links cannot expose
                # an incomplete/new set until its opportunity pointer is saved.
                save_immutable(links)
                current = call('/read', {'ids': [row['id'] for row in rows]})
                if current['records'] != head['records']:
                    continue
                (args.output / f'before-notices-{offset}-{attempt}.json').write_bytes(encoded(baseline))
                result = call('/commit', {'revision': current['revision'], 'records': rows})
                if result.get('conflict'):
                    continue
                saved = {row['id']: row for row in call('/read', {'ids': [row['id'] for row in rows]})['records']}
                if any(saved[row['id']]['data'] != row['data'] for row in rows):
                    raise ValueError('Opportunity link-set read-back differs.')
            plan['links'] += len(links)
            plan['linkedOpportunities'] += sum(bool(row['data'].get('classificationCodes')) for row in rows)
            unresolved.update(item['sourceLabel'] for row in rows for item in row['data'].get('commodityClassificationUnresolved', []))
            break
        else:
            raise ValueError('Concurrent opportunity edits prevented a safe merge.')
    plan['unresolvedLabels'] = sorted(unresolved)
    if args.apply:
        commit(entries=[{'key': ACTIVE, 'value': {'schemaVersion': 1, 'releaseId': RELEASE}}, {'key': 'procurement:taxonomy:audit', 'value': plan}])
        saved = call('/workspace', {'keys': [ACTIVE, 'procurement:taxonomy:audit']})
        if next(entry['value'] for entry in saved['entries'] if entry['key'] == 'procurement:taxonomy:audit') != plan:
            raise ValueError('Final taxonomy audit read-back failed.')
    (args.output / ('applied.json' if args.apply else 'plan.json')).write_text(json.dumps(plan, indent=2))
    print(json.dumps(plan, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base', default='https://zoer.k8s.overtheedgepaper.ca')
    parser.add_argument('--workbook', type=pathlib.Path, required=True)
    parser.add_argument('--categories', type=pathlib.Path)
    parser.add_argument('--commodity-details', type=pathlib.Path)
    parser.add_argument('--opportunities', type=pathlib.Path, help='Offline original snapshot for dry-run validation only.')
    parser.add_argument('--output', type=pathlib.Path, required=True)
    parser.add_argument('--apply', action='store_true')
    run(parser.parse_args())
